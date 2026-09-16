"""
test_qa_2026_09_16.py — non-régression du rapport QA du 16/09/2026 (mode /tmp,
sans PostgreSQL). Les preuves qui exigent une vraie base (cloisonnement admin,
persistance des snapshots, inbox de revue, anti-rejeu TOTP en base) vivent dans
test_qa_2026_09_16_db.py (job CI `migration-tests`).
"""

from __future__ import annotations

import io
import json
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from jose import jwt
from openpyxl import load_workbook

from services import auth_service
from services.auth_service import AuthUser

API_ROOT = Path(__file__).resolve().parent.parent


def auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _sign(payload: dict) -> str:
    return jwt.encode(payload, auth_service._JWT_SECRET, algorithm="HS256")


def _exp(minutes: int = 5) -> datetime:
    return datetime.now(timezone.utc) + timedelta(minutes=minutes)


# ---------------------------------------------------------------------------
# B-01 — configuration Vercel (routage FastAPI zéro-config)
# ---------------------------------------------------------------------------

class TestVercelConfig:
    def test_no_rewrite_to_a_single_function(self) -> None:
        """Depuis @vercel/python 14, une réécriture interne transmet le chemin
        RÉÉCRIT à l'application : `/(.*)` → `/api/index` faisait voir
        `/api/index` à FastAPI, d'où un 404 sur toutes les routes."""
        config = json.loads((API_ROOT / "vercel.json").read_text(encoding="utf-8"))
        assert "rewrites" not in config
        assert "routes" not in config
        assert set(config.get("functions", {})) <= {"main.py"}

    def test_main_is_the_resolved_entrypoint(self) -> None:
        """Ordre de résolution Vercel : app.py, index.py, server.py, main.py…
        Aucun fichier ne doit masquer main.py, et aucun dossier api/ ne doit
        recréer de fonction « fichier »."""
        for name in ("app.py", "index.py", "server.py"):
            for base in (API_ROOT, API_ROOT / "src", API_ROOT / "app"):
                assert not (base / name).exists(), f"{base / name} masquerait main.py"
        assert not list((API_ROOT / "api").glob("*.py")) if (API_ROOT / "api").exists() else True
        assert "app = FastAPI(" in (API_ROOT / "main.py").read_text(encoding="utf-8")

    def test_runtime_directories_are_bundled(self) -> None:
        config = json.loads((API_ROOT / "vercel.json").read_text(encoding="utf-8"))
        excluded = config["functions"]["main.py"].get("excludeFiles", "")
        for runtime_dir in ("routers", "services", "models", "utils", "data", "db",
                            "middleware", "demo", "jobs"):
            assert f"{runtime_dir}/" not in excluded, runtime_dir

    def test_health_is_served_at_its_real_path(self, client: TestClient) -> None:
        assert client.get("/health").status_code == 200
        assert client.get("/api/index").status_code == 404


# ---------------------------------------------------------------------------
# B-02 / m-01 — seul un jeton d'ACCÈS ouvre l'API
# ---------------------------------------------------------------------------

class TestAccessTokenScope:
    def test_pre_auth_token_is_rejected(self, client: TestClient) -> None:
        pre = auth_service.create_pre_auth_token(AuthUser(email="admin@carbonco.fr", role="admin", company_id=1))
        assert client.get("/auth/me", headers=auth(pre)).status_code == 401
        assert client.get("/admin/users", headers=auth(pre)).status_code == 401

    def test_front_demo_session_token_is_rejected(self, client: TestClient) -> None:
        """La session démo du front est signée avec le même secret (scope
        « demo ») : elle ne doit jamais valoir un jeton d'accès API."""
        token = _sign({"sub": "demo-session@exemplia-industrie.invalid", "role": "viewer",
                       "cid": 0, "demo": True, "scope": "demo", "exp": _exp()})
        assert client.get("/auth/me", headers=auth(token)).status_code == 401

    def test_token_without_scope_is_rejected(self, client: TestClient) -> None:
        token = _sign({"sub": "admin@carbonco.fr", "role": "admin", "cid": 1, "exp": _exp()})
        assert client.get("/auth/me", headers=auth(token)).status_code == 401

    def test_token_without_exp_is_rejected(self, client: TestClient) -> None:
        token = _sign({"sub": "admin@carbonco.fr", "role": "admin", "cid": 1, "scope": "access"})
        assert client.get("/auth/me", headers=auth(token)).status_code == 401

    @pytest.mark.parametrize("claims", [
        {"role": "superadmin", "cid": 1},
        {"role": "admin", "cid": "1"},
        {"role": "admin", "cid": True},
        {"role": "admin"},
    ])
    def test_malformed_claims_are_rejected(self, claims: dict) -> None:
        token = _sign({"sub": "x@y.fr", "scope": "access", "exp": _exp(), **claims})
        assert auth_service.decode_token(token) is None

    def test_access_token_carries_scope_and_uid(self) -> None:
        token, _ = auth_service.create_access_token(
            AuthUser(email="a@b.fr", role="analyst", company_id=7, user_id=42),
        )
        user = auth_service.decode_token(token)
        assert user is not None and user.user_id == 42 and user.company_id == 7

    def test_me_exposes_platform_admin_flag(self, client: TestClient, admin_token: str,
                                           monkeypatch) -> None:
        assert client.get("/auth/me", headers=auth(admin_token)).json()["platformAdmin"] is False
        monkeypatch.setenv("PLATFORM_ADMIN_EMAILS", "Admin@CarbonCo.fr")
        assert client.get("/auth/me", headers=auth(admin_token)).json()["platformAdmin"] is True


# ---------------------------------------------------------------------------
# B-03 — comptes de développement et administrateur de plateforme
# ---------------------------------------------------------------------------

class TestDevAccounts:
    @pytest.mark.parametrize("vercel_env", ["production", "preview"])
    def test_dev_accounts_never_open_a_session_when_hosted(self, monkeypatch, vercel_env) -> None:
        monkeypatch.setenv("VERCEL_ENV", vercel_env)
        assert auth_service.authenticate("admin@carbonco.fr", "Admin2024!") is None
        assert auth_service.authenticate("demo@carbonco.fr", "CarbonCo2024!") is None

    def test_dev_accounts_work_in_local_development(self, monkeypatch) -> None:
        monkeypatch.delenv("VERCEL_ENV", raising=False)
        user = auth_service.authenticate("viewer@carbonco.fr", "Viewer2024!")
        assert user is not None and user.role == "viewer"

    def test_public_passwords_are_refused_in_production_even_with_a_db(self, monkeypatch) -> None:
        """Un compte ensemencé avant le correctif avec un mot de passe public
        ne doit plus ouvrir de session : refus AVANT toute lecture en base."""
        monkeypatch.setenv("VERCEL_ENV", "production")
        monkeypatch.setattr(auth_service, "db_available", lambda: True)

        def _must_not_read(_email):
            raise AssertionError("la base ne doit pas être consultée")

        monkeypatch.setattr(auth_service, "_get_user_from_db", _must_not_read)
        assert auth_service.authenticate("admin@carbonco.fr", "Admin2024!") is None

    def test_seeding_requires_explicit_opt_in(self, monkeypatch) -> None:
        monkeypatch.delenv("VERCEL_ENV", raising=False)
        monkeypatch.delenv("CARBONCO_SEED_DEV_ACCOUNTS", raising=False)
        monkeypatch.setattr(auth_service, "_DEMO_USERS_SEEDED", False)
        monkeypatch.setattr(auth_service, "db_available", lambda: True)

        def _must_not_connect(*_a, **_k):
            raise AssertionError("aucun ensemencement sans opt-in")

        monkeypatch.setattr(auth_service, "get_db", _must_not_connect)
        auth_service._ensure_default_users()

    def test_seeding_never_happens_in_production(self, monkeypatch) -> None:
        monkeypatch.setenv("VERCEL_ENV", "production")
        monkeypatch.setenv("CARBONCO_SEED_DEV_ACCOUNTS", "1")
        monkeypatch.setattr(auth_service, "_DEMO_USERS_SEEDED", False)
        monkeypatch.setattr(auth_service, "db_available", lambda: True)
        monkeypatch.setattr(auth_service, "get_db", lambda *a, **k: (_ for _ in ()).throw(AssertionError()))
        auth_service._ensure_default_users()


class TestPlatformAdmin:
    def test_requires_admin_role_listed_email_and_no_demo(self, monkeypatch) -> None:
        monkeypatch.setenv("PLATFORM_ADMIN_EMAILS", " ops@carbonco.fr , boss@carbonco.fr")
        admin = AuthUser(email="ops@carbonco.fr", role="admin", company_id=1)
        assert auth_service.is_platform_admin(admin)
        assert not auth_service.is_platform_admin(admin.model_copy(update={"role": "analyst"}))
        assert not auth_service.is_platform_admin(admin.model_copy(update={"is_demo": True}))
        assert not auth_service.is_platform_admin(admin.model_copy(update={"email": "x@carbonco.fr"}))

    def test_empty_list_grants_nobody(self, monkeypatch) -> None:
        monkeypatch.delenv("PLATFORM_ADMIN_EMAILS", raising=False)
        assert not auth_service.is_platform_admin(AuthUser(email="admin@carbonco.fr", role="admin"))

    def test_target_company_scoping(self, monkeypatch) -> None:
        from fastapi import HTTPException

        from routers.admin import _target_company

        admin = AuthUser(email="a@org1.fr", role="admin", company_id=1)
        assert _target_company(admin, None) == 1
        assert _target_company(admin, 1) == 1
        with pytest.raises(HTTPException) as exc:
            _target_company(admin, 2)
        assert exc.value.status_code == 403
        monkeypatch.setenv("PLATFORM_ADMIN_EMAILS", "a@org1.fr")
        assert _target_company(admin, 2) == 2

    @pytest.mark.parametrize("password,ok", [
        ("Admin2024!", False),                      # mot de passe public
        ("court1A!", False),
        ("Motdepasse123!", True),                   # 12+ car., 4 types
        ("MotDePasseLong14", True),                 # 14+ car., 3 types
        ("motdepasselongsanschiffre", False),
        ("cheval correct pile agrafe lune verte soleil", True),  # 7 mots
    ])
    def test_password_policy_follows_cnil_2022_100(self, password: str, ok: bool) -> None:
        from routers.admin import _password_problem

        assert (_password_problem(password) is None) is ok

    def test_admin_routes_require_admin(self, client: TestClient, analyst_token: str) -> None:
        assert client.get("/admin/users").status_code == 401
        assert client.get("/admin/users", headers=auth(analyst_token)).status_code == 403


# ---------------------------------------------------------------------------
# Accès tenant : plus de repli anonyme sur l'organisation n°1
# ---------------------------------------------------------------------------

TENANT_GETS = [
    "/carbon/snapshot", "/vsme/snapshot", "/esg/snapshot", "/finance/snapshot",
    "/dashboard/consolidated", "/suppliers", "/suppliers/scope3",
    "/materialite/positions", "/materialite/assessments", "/beges/filings",
    "/strategic-mapping/adhesion-volontaire", "/history/carbon", "/alerts/rules",
    "/alerts/history", "/ingest/status", "/dpp/products", "/copilot/tools",
]


class TestTenantEndpointsRequireToken:
    @pytest.mark.parametrize("path", TENANT_GETS)
    def test_get_without_token_is_401(self, client: TestClient, path: str) -> None:
        assert client.get(path).status_code == 401

    @pytest.mark.parametrize("path", ["/ingest", "/report/generate", "/alerts/evaluate"])
    def test_post_without_token_is_401(self, client: TestClient, path: str) -> None:
        assert client.post(path).status_code == 401

    def test_invalid_token_is_401(self, client: TestClient) -> None:
        assert client.get("/carbon/snapshot", headers=auth("not-a-jwt")).status_code == 401


class TestSnapshotsNeverServeDemoDataToATenant:
    def test_db_mode_without_import_is_404_no_snapshot(self, client: TestClient,
                                                      analyst_token: str, monkeypatch) -> None:
        from routers import _snapshots

        monkeypatch.setattr(_snapshots, "db_available", lambda: True)
        monkeypatch.setattr(_snapshots, "read_snapshot", lambda *a, **k: None)
        for domain in ("carbon", "vsme", "esg", "finance"):
            resp = client.get(f"/{domain}/snapshot", headers=auth(analyst_token))
            assert resp.status_code == 404
            assert resp.json()["detail"]["error"] == "no_snapshot"

    def test_store_failure_is_503(self, client: TestClient, analyst_token: str, monkeypatch) -> None:
        from routers import _snapshots
        from services.snapshot_cache import SnapshotStoreError

        def _boom(*_a, **_k):
            raise SnapshotStoreError("Lecture du snapshot carbon impossible.")

        monkeypatch.setattr(_snapshots, "read_snapshot", _boom)
        assert client.get("/carbon/snapshot", headers=auth(analyst_token)).status_code == 503

    def test_store_failure_on_dashboard_is_503(self, client: TestClient, analyst_token: str,
                                               monkeypatch) -> None:
        from services import aggregation_service
        from services.snapshot_cache import SnapshotStoreError

        def _boom(*_a, **_k):
            raise SnapshotStoreError("Lecture de l'état des snapshots impossible.")

        monkeypatch.setattr(aggregation_service, "cache_status", _boom)
        resp = client.get("/dashboard/consolidated", headers=auth(analyst_token))
        assert resp.status_code == 503
        assert "snapshots" in resp.json()["detail"]

    def test_master_ingest_forbidden_in_production(self, client: TestClient,
                                                  admin_token: str, monkeypatch) -> None:
        monkeypatch.setenv("VERCEL_ENV", "production")
        monkeypatch.delenv("PLATFORM_ADMIN_EMAILS", raising=False)
        assert client.post("/ingest", headers=auth(admin_token)).status_code == 403

    def test_master_ingest_requires_admin(self, client: TestClient, analyst_token: str) -> None:
        assert client.post("/ingest", headers=auth(analyst_token)).status_code == 403


# ---------------------------------------------------------------------------
# B-04 — persistance des snapshots
# ---------------------------------------------------------------------------

class _FakeCursor:
    """Curseur façon RealDictCursor : lignes = dict (row[0] lèverait KeyError)."""

    def __init__(self, log: list[str]) -> None:
        self.log = log
        self._rows: list[dict] = []

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def execute(self, sql: str, params=None) -> None:
        self.log.append(" ".join(sql.split()))
        if "next_version" in sql:
            self._rows = [{"next_version": 3}]
        elif sql.lstrip().startswith("INSERT INTO snapshots"):
            self._rows = [{"id": 99, "generated_at": datetime(2026, 9, 16, tzinfo=timezone.utc)}]
        else:
            self._rows = []

    def fetchone(self):
        return self._rows[0] if self._rows else None


class _FakeConn:
    def __init__(self, log: list[str]) -> None:
        self.log = log

    def cursor(self):
        return _FakeCursor(self.log)


class TestSnapshotPersistence:
    def test_pg_write_reads_dict_rows(self, monkeypatch) -> None:
        from services import snapshot_cache

        log: list[str] = []
        seen_company: list[int | None] = []

        @contextmanager
        def fake_get_db(company_id=None):
            seen_company.append(company_id)
            yield _FakeConn(log)

        monkeypatch.setattr(snapshot_cache, "db_available", lambda: True)
        monkeypatch.setattr(snapshot_cache, "get_db", fake_get_db)
        result = snapshot_cache.write_snapshot("carbon", {"k": 1}, company_id=5, source="user_upload")
        assert result == {"id": 99, "version": 3, "generatedAt": "2026-09-16T00:00:00+00:00",
                          "source": "user_upload"}
        assert seen_company == [5], "l'écriture doit poser le contexte tenant (RLS)"
        assert any("pg_advisory_xact_lock" in q for q in log)

    def test_pg_write_failure_is_never_a_silent_success(self, monkeypatch, tmp_path) -> None:
        from services import snapshot_cache

        @contextmanager
        def broken_get_db(company_id=None):
            raise RuntimeError("connexion refusée")
            yield  # pragma: no cover

        monkeypatch.setenv("CARBONCO_CACHE_DIR", str(tmp_path))
        monkeypatch.setattr(snapshot_cache, "db_available", lambda: True)
        monkeypatch.setattr(snapshot_cache, "get_db", broken_get_db)
        with pytest.raises(snapshot_cache.SnapshotStoreError):
            snapshot_cache.write_snapshot("carbon", {"k": 1}, company_id=5)
        assert not list(tmp_path.rglob("*.json")), "aucun repli fichier global"
        with pytest.raises(snapshot_cache.SnapshotStoreError):
            snapshot_cache.read_snapshot("carbon", company_id=5)

    def test_file_mode_is_isolated_per_company(self, monkeypatch, tmp_path) -> None:
        from services import snapshot_cache

        monkeypatch.setenv("CARBONCO_CACHE_DIR", str(tmp_path))
        snapshot_cache.write_snapshot("carbon", {"owner": 1}, company_id=1)
        assert snapshot_cache.read_snapshot("carbon", company_id=1) == {"owner": 1}
        assert snapshot_cache.read_snapshot("carbon", company_id=2) is None
        assert snapshot_cache.cache_status(company_id=2)["carbon"] == {"exists": False}
        snapshot_cache.invalidate("carbon", company_id=2)
        assert snapshot_cache.read_snapshot("carbon", company_id=1) == {"owner": 1}

    def test_unknown_domain_is_rejected(self, monkeypatch, tmp_path) -> None:
        from services import snapshot_cache

        monkeypatch.setenv("CARBONCO_CACHE_DIR", str(tmp_path))
        with pytest.raises(ValueError):
            snapshot_cache.write_snapshot("../../etc", {}, company_id=1)

    def test_ingest_uploaded_reports_unsaved_import(self, client: TestClient, analyst_token: str,
                                                   monkeypatch) -> None:
        from routers import excel as excel_router
        from services.snapshot_cache import SnapshotStoreError

        monkeypatch.setattr(excel_router, "build_carbon_snapshot_from_bytes", lambda *a, **k: {
            "generatedAt": "2026-09-16T00:00:00Z", "carbon": {},
            "validation": {"status": "ok", "failures": [], "warnings": []},
        })
        monkeypatch.setattr(excel_router, "write_snapshot",
                            lambda *a, **k: (_ for _ in ()).throw(SnapshotStoreError("x")))
        template = (API_ROOT / "data" / "CarbonCo_Calcul_Carbone_v2.xlsx").read_bytes()
        resp = client.post(
            "/excel/ingest-uploaded",
            files={"file": ("bilan.xlsx", template,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"domain": "carbon"},
            headers=auth(analyst_token),
        )
        assert resp.status_code == 503
        assert resp.json()["detail"]["error"] == "snapshot_not_saved"


# ---------------------------------------------------------------------------
# B-07 — évaluation des alertes par le cron
# ---------------------------------------------------------------------------

SERVICE_TOKEN = "cron-service-token-0123456789abcdef"


@pytest.fixture()
def two_tenant_rules(monkeypatch):
    from routers import alerts

    monkeypatch.setattr(alerts, "_MEM_RULES", [])
    monkeypatch.setattr(alerts, "_MEM_NOTIFS", [])
    monkeypatch.setattr(alerts, "_MEM_HISTORY", alerts.deque(maxlen=100))
    base = {"domain": "carbon", "field_path": "carbon.totalS123Tco2e", "operator": "gt",
            "threshold": None, "mode": "missing", "channel": "inapp", "destination": None,
            "is_active": True, "last_fired_at": None, "created_at": "2026-09-16T00:00:00+00:00"}
    alerts._MEM_RULES.extend([
        {**base, "id": 1, "company_id": 1, "name": "Org 1"},
        {**base, "id": 2, "company_id": 2, "name": "Org 2"},
    ])
    return alerts


class TestCronEvaluation:
    def test_service_token_evaluates_every_tenant(self, client: TestClient, monkeypatch,
                                                  two_tenant_rules) -> None:
        monkeypatch.setenv("CRON_SERVICE_TOKEN", SERVICE_TOKEN)
        resp = client.post("/alerts/evaluate", headers=auth(SERVICE_TOKEN))
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["evaluated"] == 2 and data["companies"] == 2
        assert {a["company_id"] for a in data["alerts"]} == {1, 2}

    def test_duplicate_cron_delivery_does_not_refire(self, client: TestClient, monkeypatch,
                                                     two_tenant_rules) -> None:
        monkeypatch.setenv("CRON_SERVICE_TOKEN", SERVICE_TOKEN)
        first = client.post("/alerts/evaluate", headers=auth(SERVICE_TOKEN)).json()
        second = client.post("/alerts/evaluate", headers=auth(SERVICE_TOKEN)).json()
        assert first["fired"] == 2
        assert second["fired"] == 0

    def test_analyst_only_evaluates_own_tenant(self, client: TestClient, analyst_token: str,
                                               two_tenant_rules) -> None:
        data = client.post("/alerts/evaluate", headers=auth(analyst_token)).json()
        assert data["evaluated"] == 1
        history = client.get("/alerts/history", headers=auth(analyst_token)).json()
        assert all(a["rule_id"] == 1 for a in history["alerts"])

    @pytest.mark.parametrize("configured,presented", [
        (SERVICE_TOKEN, "wrong-token-0123456789abcdef"),
        ("", SERVICE_TOKEN),
        ("short", "short"),
    ])
    def test_bad_or_weak_service_tokens_are_refused(self, client: TestClient, monkeypatch,
                                                    configured: str, presented: str) -> None:
        monkeypatch.setenv("CRON_SERVICE_TOKEN", configured)
        assert client.post("/alerts/evaluate", headers=auth(presented)).status_code == 401

    def test_reminders_are_scoped_for_users(self, client: TestClient, analyst_token: str,
                                            monkeypatch) -> None:
        from services import beges_filings_service, supplier_campaigns_service

        seen: dict[str, object] = {}

        def _recorder(key: str):
            def run(company_ids=None):
                seen[key] = company_ids
                return {"checked": 0}
            return run

        monkeypatch.setattr(beges_filings_service, "run_reminders", _recorder("beges"))
        monkeypatch.setattr(supplier_campaigns_service, "run_campaign_reminders", _recorder("suppliers"))
        client.post("/beges/reminders/run", headers=auth(analyst_token))
        client.post("/suppliers/campaigns/reminders/run", headers=auth(analyst_token))
        assert seen == {"beges": {1}, "suppliers": {1}}

    def test_reminders_run_for_all_tenants_with_service_token(self, client: TestClient,
                                                              monkeypatch) -> None:
        from services import beges_filings_service

        monkeypatch.setenv("CRON_SERVICE_TOKEN", SERVICE_TOKEN)
        calls: list[object] = []
        monkeypatch.setattr(beges_filings_service, "run_reminders",
                            lambda company_ids=None: calls.append(company_ids) or {"checked": 0})
        assert client.post("/beges/reminders/run", headers=auth(SERVICE_TOKEN)).status_code == 200
        assert calls == [None]


# ---------------------------------------------------------------------------
# Moteur de formules : plus d'eval() sur les classeurs importés
# ---------------------------------------------------------------------------

class TestSafeFormula:
    @pytest.mark.parametrize("expr,expected", [
        ("1+2*3", 7), ("(1+2)*3", 9), ("2**3", 8), ("-4+1", -3), ("10/4", 2.5),
        ("3>2", True), ("1==1.0", True), ("'a'=='a'", True), ("not False", True),
        ("True and False", False), ("0 or 5", 5), ("1<2<3", True),
    ])
    def test_arithmetic_and_logic(self, expr: str, expected) -> None:
        from utils.safe_formula import safe_eval

        assert safe_eval(expr) == expected

    @pytest.mark.parametrize("expr", [
        "().__class__.__base__.__subclasses__()",
        "__import__('os').system('echo pwned')",
        "open('x')", "[1, 2][0]", "x", "lambda: 1", "[c for c in 'ab']",
        "10 % 3", "9**9**9", "1/0", "1e308*10", "'a' < 1", "{1: 2}",
    ])
    def test_everything_else_is_refused(self, expr: str) -> None:
        from utils.safe_formula import UnsafeFormulaError, safe_eval

        with pytest.raises(UnsafeFormulaError):
            safe_eval(expr)

    def test_carbon_engine_blocks_a_crafted_workbook(self) -> None:
        from openpyxl import Workbook

        from services import carbon_service

        wb = Workbook()
        ws = wb.active
        ws.title = "Synthese_GES"
        ws["C10"] = "=().__class__.__base__.__subclasses__().__len__()"
        with pytest.raises(carbon_service.CarbonServiceError):
            carbon_service._evaluate_formula_cell(wb, "Synthese_GES", "C10", {}, set())

    def test_if_branches_are_evaluated(self) -> None:
        from openpyxl import Workbook

        from services import carbon_service

        wb = Workbook()
        ws = wb.active
        ws.title = "S"
        ws["A1"], ws["A2"] = 3, 4
        ws["B1"] = "=IF(A1>2,A2/2,0)"
        ws["B2"] = '=IF(A1>5,"haut","bas")'
        assert carbon_service._evaluate_formula_cell(wb, "S", "B1", {}, set()) == 2.0
        assert carbon_service._evaluate_formula_cell(wb, "S", "B2", {}, set()) == "bas"


# ---------------------------------------------------------------------------
# M-02 — agrégats recalculés depuis les lignes de détail
# ---------------------------------------------------------------------------

def _reader(cells: dict[tuple[str, str], object]):
    return lambda sheet, ref: cells.get((sheet, ref))


class TestWorkbookTotals:
    def test_ghg_energy_taxonomy_cbam(self) -> None:
        from services.carbon import workbook_totals as wt

        cells = {
            ("Synthese_GES", "C6"): 20.5, ("Synthese_GES", "C13"): 4.0,
            ("Synthese_GES", "C14"): 1.0, ("Synthese_GES", "C16"): 2.0,
            ("Synthese_GES", "C22"): 2.2, ("Synthese_GES", "C20"): 12.3,
            ("Energie", "E10"): 100, ("Energie", "E16"): 50,
            # ligne 14 alignée (objectif renseigné), ligne 15 sans objectif → exclue
            **{("Taxonomie", f"{c}14"): v for c, v in
               {"C": "O", "D": 1, "E": "O", "F": "O", "H": 2_000_000, "I": 100_000, "J": 0}.items()},
            **{("Taxonomie", f"{c}15"): v for c, v in
               {"C": "O", "E": "O", "F": "O", "H": 1_000_000}.items()},
            ("CBAM", "B4"): 70, ("CBAM", "B5"): 2026,
            ("CBAM", "B9"): "Acier et fer", ("CBAM", "I9"): 185, ("CBAM", "J9"): 0,
        }
        res = wt.compute_totals(_reader(cells), revenue_eur=10_000_000, fte=50,
                                capex_eur=1_000_000, opex_eur=200_000)
        v = res.values
        assert res.failures == []
        assert v["carbon.scope1Tco2e"] == 20.5
        assert v["carbon.scope2LbTco2e"] == 5.0          # électricité LB + chaleur
        assert v["carbon.scope2MbTco2e"] == 3.0          # électricité MB + chaleur
        assert v["carbon.scope3Tco2e"] == 14.5
        assert v["carbon.totalS123Tco2e"] == 40.0        # S1 + S2 LB + S3
        assert v["carbon.intensityRevenueTco2ePerMEur"] == 4.0
        assert v["carbon.intensityFteTco2ePerFte"] == 0.8
        assert (v["carbon.shareScope1Pct"], v["carbon.shareScope2Pct"],
                v["carbon.shareScope3Pct"]) == (51.2, 12.5, 36.2)
        assert v["energy.consumptionMWh"] == 150 and v["energy.renewableSharePct"] == 33.3
        assert v["taxonomy.turnoverAlignedPct"] == 20.0
        assert v["taxonomy.capexAlignedPct"] == 10.0
        assert v["taxonomy.opexAlignedPct"] == 0.0
        assert v["cbam.estimatedCostEur"] == 323.75       # 185 t × 70 € × 2,5 %

    @pytest.mark.parametrize("year,share", [
        (2025, 0.0), (2026, 0.025), (2027, 0.05), (2028, 0.10), (2029, 0.225),
        (2030, 0.485), (2031, 0.61), (2032, 0.735), (2033, 0.86), (2034, 1.0), (2040, 1.0),
        (None, None),
    ])
    def test_cbam_payable_share_follows_article_31(self, year, share) -> None:
        from services.carbon.workbook_totals import cbam_payable_share

        assert cbam_payable_share(year) == share

    def test_cbam_legacy_quarter_label_gives_the_year(self) -> None:
        from services.carbon import workbook_totals as wt

        cells = {("CBAM", "B4"): 80, ("CBAM", "B5"): "T1 2027",
                 ("CBAM", "B9"): "Aluminium", ("CBAM", "I9"): 10, ("CBAM", "J9"): 0}
        res = wt.compute_totals(_reader(cells), revenue_eur=None, fte=None, capex_eur=None, opex_eur=None)
        assert res.values["cbam.estimatedCostEur"] == 40.0  # 10 × 80 × 5 %

    def test_excel_error_is_a_failure_not_a_number(self) -> None:
        from services.carbon import workbook_totals as wt

        res = wt.compute_totals(_reader({("Synthese_GES", "C20"): "#N/A"}),
                                revenue_eur=1, fte=1, capex_eur=1, opex_eur=1)
        assert "carbon.totalS123Tco2e" not in res.values
        assert any("#N/A" in f for f in res.failures)

    def test_zero_denominators_give_no_ratio(self) -> None:
        from services.carbon import workbook_totals as wt

        res = wt.compute_totals(_reader({}), revenue_eur=0, fte=None, capex_eur=None, opex_eur=None)
        assert res.values["carbon.totalS123Tco2e"] == 0
        assert res.values["carbon.shareScope1Pct"] is None
        assert res.values["carbon.intensityRevenueTco2ePerMEur"] is None
        assert res.values["energy.renewableSharePct"] is None
        assert res.values["taxonomy.turnoverAlignedPct"] is None

    def test_differs_tolerance(self) -> None:
        from services.carbon.workbook_totals import differs

        assert differs(5955, 22.7)
        assert not differs(22.7, 22.7004)
        assert not differs(None, 22.7)
        assert differs("n/a", 1.0)


def _template_with(cells: dict[tuple[str, str], object]) -> bytes:
    wb = load_workbook(API_ROOT / "data" / "CarbonCo_Calcul_Carbone_v2.xlsx")
    for (sheet, ref), value in cells.items():
        wb[sheet][ref] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


PARAMS = {("Paramètres", "B4"): "Test SAS", ("Paramètres", "B7"): 2025,
          ("Paramètres", "B9"): 10_000_000, ("Paramètres", "B11"): 50}


class TestCarbonImportFromTemplate:
    def test_official_template_has_no_frozen_totals_nor_demo_company(self) -> None:
        wb = load_workbook(API_ROOT / "data" / "CarbonCo_Calcul_Carbone_v2.xlsx")
        synth = wb["Synthese_GES"]
        for ref in ("C10", "C15", "C17", "C35", "C47", "C50", "C51", "C53", "C54", "C55"):
            assert str(synth[ref].value).startswith("="), f"Synthese_GES!{ref} figé"
        for sheet, ref in (("Energie", "E19"), ("Energie", "E20"), ("Taxonomie", "E27"),
                           ("CBAM", "M24"), ("CBAM", "B6")):
            assert str(wb[sheet][ref].value).startswith("="), f"{sheet}!{ref} figé"
        assert wb["Paramètres"]["B4"].value is None
        assert wb["Paramètres"]["B9"].value is None

    def test_untouched_template_is_rejected(self) -> None:
        from services.carbon_service import build_carbon_snapshot_from_bytes

        template = (API_ROOT / "data" / "CarbonCo_Calcul_Carbone_v2.xlsx").read_bytes()
        snap = build_carbon_snapshot_from_bytes(template)
        assert snap["validation"]["status"] == "failed"
        failures = " ".join(snap["validation"]["failures"])
        assert "Aucune émission calculée" in failures
        assert "chiffre d'affaires net" in failures

    def test_frozen_totals_of_an_old_template_are_ignored(self) -> None:
        """Ancien modèle : totaux figés (5 955) mais détail réel (20,5 + 2,2)."""
        from services.carbon_service import build_carbon_snapshot_from_bytes

        # Classeur « enregistré par un tableur » : chaque ligne de détail porte
        # sa valeur calculée (openpyxl ne conserve pas les résultats de calcul).
        detail = {("Synthese_GES", ref): 0 for ref in (
            "C6", "C7", "C8", "C9", "C13", "C14", "C16", *(f"C{r}" for r in range(20, 35)),
        )}
        content = _template_with({
            **PARAMS, **detail,
            ("Synthese_GES", "C6"): 20.5, ("Synthese_GES", "C22"): 2.2,
            ("Synthese_GES", "C10"): 1336, ("Synthese_GES", "C35"): 3685,
            ("Synthese_GES", "C47"): 5955,
            ("Energie", "E10"): 100, ("Energie", "E16"): 50,
            **{("CBAM", f"I{r}"): 0 for r in range(9, 24)},
        })
        snap = build_carbon_snapshot_from_bytes(content)
        assert snap["carbon"]["totalS123Tco2e"] == 22.7
        assert snap["carbon"]["scope1Tco2e"] == 20.5
        assert snap["carbon"]["intensityRevenueTco2ePerMEur"] == 2.27
        assert snap["validation"]["status"] == "warning"
        assert any("ne correspondent pas au détail" in w for w in snap["validation"]["warnings"])

    def test_workbook_saved_without_results_is_rejected_with_guidance(self) -> None:
        from services.carbon_service import build_carbon_snapshot_from_bytes

        content = _template_with({**PARAMS, ("Scope_1", "C6"): "Gaz naturel (PCI)",
                                  ("Scope_1", "D6"): 100000})
        snap = build_carbon_snapshot_from_bytes(content)
        assert snap["validation"]["status"] == "failed"
        assert snap["carbon"]["totalS123Tco2e"] is None
        assert any("enregistrez-le" in f for f in snap["validation"]["failures"])


# ---------------------------------------------------------------------------
# M-01 — complétude VSME honnête
# ---------------------------------------------------------------------------

class TestVsmeCompleteness:
    def test_labels_are_not_values(self) -> None:
        from services import vsme_mapping_service as vm

        snapshot = {"profile": {"etp": "Effectifs (ETP)", "caNet": "Chiffre d'affaires net (€)"},
                    "environnement": {"scope1Tco2e": "12,5"}}
        rows = {r["code"]: r for r in vm.map_datapoints(snapshot, {})}
        assert rows["B1-3"]["status"] == "missing"
        by_path = {dp["snapshot"]: dp["code"] for dp in vm.vsme_catalog.all_datapoints() if dp.get("snapshot")}
        assert rows[by_path["environnement.scope1Tco2e"]]["status"] == "auto"

    def test_empty_organisation_is_not_complete(self) -> None:
        from services import vsme_mapping_service as vm

        comp = vm.completeness(vm.map_datapoints({}, {}))
        assert comp["mandatory_filled"] == 0 and comp["overall_pct"] == 0

    def test_master_esg_workbook_points_to_values(self) -> None:
        wb = load_workbook(API_ROOT / "data" / "CarbonCo_ESG_Social.xlsx")
        for name in wb.defined_names:
            if name.startswith("CC_VSME_"):
                assert "!$E$" in wb.defined_names[name].attr_text, name

    def test_db_mode_never_uses_the_master_workbook(self, monkeypatch) -> None:
        from services import snapshot_cache
        from services import vsme_mapping_service as vm

        monkeypatch.setattr(vm, "db_available", lambda: True)
        monkeypatch.setattr(snapshot_cache, "read_snapshot", lambda *a, **k: None)
        assert vm._snapshot_dict(3) == {}


# ---------------------------------------------------------------------------
# M-06 — santé
# ---------------------------------------------------------------------------

class TestHealth:
    @pytest.mark.parametrize("db,storage,prod,expected", [
        ("ok", "ok", True, "ok"),
        ("down", "ok", False, "degraded"),
        ("ok", "down", False, "degraded"),
        ("not_configured", "local", False, "ok"),
        ("not_configured", "local", True, "degraded"),
    ])
    def test_overall_status(self, monkeypatch, db, storage, prod, expected) -> None:
        from routers import health

        monkeypatch.setattr(health, "is_production", lambda: prod)
        assert health.overall_status(db, storage) == expected

    def test_db_down_is_reported(self, client: TestClient, monkeypatch) -> None:
        from routers import health

        monkeypatch.setattr(health, "_db_status", lambda: "down")
        body = client.get("/health").json()
        assert body["status"] == "degraded" and body["db"] == "down"


# ---------------------------------------------------------------------------
# M-15 — CORS et origine des requêtes de session
# ---------------------------------------------------------------------------

class TestCors:
    @pytest.mark.parametrize("origin,allowed", [
        ("https://carbon-snowy-nine.vercel.app", True),
        ("https://carbon-ezy1l3mez-ludovics-projects-159c139c.vercel.app", True),
        ("https://carbon-git-feature-x-ludovics-projects-159c139c.vercel.app", True),
        ("https://carbonevil.vercel.app", False),
        ("https://carbon-attacker.vercel.app", False),
        ("https://carbon-ezy1l3mez-otherteam.vercel.app", False),
        ("https://carbon-snowy-nine.vercel.app.evil.test", False),
        ("https://finance-platform-x.vercel.app", False),
    ])
    def test_policy(self, origin: str, allowed: bool) -> None:
        from utils.cors import is_allowed_origin

        assert is_allowed_origin(origin) is allowed

    def test_localhost_only_outside_production(self, monkeypatch) -> None:
        from utils.cors import allowed_origins

        assert "http://localhost:3003" in allowed_origins()
        monkeypatch.setenv("VERCEL_ENV", "production")
        assert not any("localhost" in o for o in allowed_origins())

    def test_preflight_headers(self, client: TestClient) -> None:
        ok = client.options("/auth/me", headers={
            "Origin": "https://carbon-snowy-nine.vercel.app",
            "Access-Control-Request-Method": "GET",
        })
        assert ok.headers.get("access-control-allow-origin") == "https://carbon-snowy-nine.vercel.app"
        ko = client.options("/auth/me", headers={
            "Origin": "https://carbon-attacker.vercel.app",
            "Access-Control-Request-Method": "GET",
        })
        assert "access-control-allow-origin" not in ko.headers

    def test_foreign_origin_cannot_rotate_the_session(self, client: TestClient) -> None:
        resp = client.post("/auth/refresh", headers={"Origin": "https://carbon-attacker.vercel.app"})
        assert resp.status_code == 403
        resp = client.post("/auth/logout", headers={"Origin": "https://carbon-attacker.vercel.app"})
        assert resp.status_code == 403


# ---------------------------------------------------------------------------
# m-02 — anti-rejeu TOTP (mode fichier) ; journal d'audit avant 044
# ---------------------------------------------------------------------------

class TestTotpReplayAndAudit:
    def test_claim_step_is_single_use(self, monkeypatch, tmp_path) -> None:
        from services import totp_service

        monkeypatch.setenv("CARBONCO_CACHE_DIR", str(tmp_path))
        assert totp_service._claim_step("u@x.fr", 100) is True
        assert totp_service._claim_step("u@x.fr", 100) is False
        assert totp_service._claim_step("v@x.fr", 100) is True

    def test_missing_044_table_degrades_gracefully(self, monkeypatch) -> None:
        from services import totp_service

        class _Undefined(Exception):
            pgcode = "42P01"

        @contextmanager
        def missing_table(*_a, **_k):
            raise _Undefined("relation user_totp_used_steps does not exist")
            yield  # pragma: no cover

        monkeypatch.setattr(totp_service, "db_available", lambda: True)
        monkeypatch.setattr(totp_service, "get_db", missing_table)
        assert totp_service._claim_step("u@x.fr", 1) is True

    def test_new_audit_types_fall_back_before_044(self, monkeypatch) -> None:
        from services import audit_service

        class _Check(Exception):
            pgcode = "23514"

        calls: list[tuple[str, dict | None]] = []

        @contextmanager
        def fake_get_db(company_id=None):
            class _Cur:
                def __enter__(self):
                    return self

                def __exit__(self, *exc):
                    return False

                def execute(self, sql, params):
                    calls.append((params[2], json.loads(params[6]) if params[6] else None))
                    if params[2] == "2fa_disable":
                        raise _Check("violates check constraint audit_eventtype_check")

                def fetchone(self):
                    return {"id": 1, "created_at": None}

            class _Conn:
                def cursor(self):
                    return _Cur()

            yield _Conn()

        monkeypatch.setattr(audit_service, "db_available", lambda: True)
        monkeypatch.setattr(audit_service, "get_db", fake_get_db)
        audit_service.log_event("2fa_disable", "2FA désactivée", user="u@x.fr", company_id=1)
        assert [c[0] for c in calls] == ["2fa_disable", "2fa_fail"]
        assert calls[1][1] == {"event_type": "2fa_disable"}


# ---------------------------------------------------------------------------
# M-14 — historique
# ---------------------------------------------------------------------------

def test_history_requires_token(client: TestClient, analyst_token: str) -> None:
    assert client.get("/history/carbon").status_code == 401
    assert client.get("/history/carbon/1").status_code == 401
    resp = client.get("/history/carbon", headers=auth(analyst_token))
    assert resp.status_code == 200 and resp.json()["available"] is False
